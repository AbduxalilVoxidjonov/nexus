"""Fayl, eslatma va Excel/CSV toollari — kengaytma moduli.

Eksport (registry kontrakti):
    TOOL_DECLARATIONS: list[dict]
    HANDLERS: dict[str, handler]   — name -> async handler(args) -> (ok, out) | dict

Xavfsizlik:
    * Faqat `~`, /tmp va /Volumes ostidagi yo'llar (ALLOWED_ROOTS) ruxsat etiladi.
    * `path_is_sensitive` (nexus.safety yoki shu yerdagi fallback) rost bo'lsa — rad.
    * `read_file` chiqishi `redact_secrets` orqali tozalanadi.
    * Mavjud faylni overwrite qilish va `delete_file` — `needs_confirmation`; registry gate
      tasdiqdan keyin handlerni `confirmed=True` bilan qayta chaqiradi.
    * `delete_file` faylni Finder orqali Savatga (Trash) ko'chiradi — yo'q qilmaydi.

Aliaslar: "desktop / ish stoli / рабочий стол" → ~/Desktop, "downloads / yuklamalar / загрузки"
→ ~/Downloads, "documents / hujjatlar / документы" → ~/Documents, "home" → ~.
"""

from __future__ import annotations

import asyncio
import csv
import json
import logging
import os
import re
import sys
from collections.abc import Callable
from datetime import datetime
from pathlib import Path
from typing import Any

from nexus.macos_actions import _as_str, run_applescript, run_shell

log = logging.getLogger("nexus.files")

# ---------------------------------------------------------------------------
# Xavfsizlik yordamchilari (nexus.safety bo'lsa undan, bo'lmasa fallback)
# ---------------------------------------------------------------------------
_FALLBACK_SENSITIVE_NAMES = frozenset(
    {
        ".zshrc",
        ".zprofile",
        ".zshenv",
        ".bashrc",
        ".bash_profile",
        ".profile",
        ".bash_login",
        ".zlogin",
        ".netrc",
        ".pgpass",
        ".env",
        "authorized_keys",
        "known_hosts",
        "id_rsa",
        "id_ed25519",
        "id_ecdsa",
        "sudoers",
        "crontab",
        "hosts",
    }
)
_FALLBACK_SENSITIVE_DIRS = (
    "/.ssh",
    "/.gnupg",
    "/.aws",
    "/.config/gcloud",
    "/.docker",
    "/.kube",
    "/library/launchagents",
    "/library/launchdaemons",
    "/library/keychains",
    "/library/preferences",
    "/library/application support/com.apple.tcc",
    "/system",
    "/etc",
    "/private/etc",
    "/usr",
    "/bin",
    "/sbin",
    "/var/db",
    "/.config/autostart",
)
_FALLBACK_SENSITIVE_SUFFIXES = (
    ".plist",
    ".command",
    ".sh",
    ".zsh",
    ".bash",
    ".terminal",
    ".workflow",
    ".app",
    ".pem",
    ".key",
    ".p12",
    ".keychain",
    ".keychain-db",
)
_SECRET_ASSIGN = re.compile(
    r"((?:API[_-]?KEY|TOKEN|SECRET|PASSWORD|PASSWD|CREDENTIAL|AUTH|BEARER|"
    r"PRIVATE[_-]?KEY|ACCESS[_-]?KEY)\w*\s*[=:]\s*[\"']?)([^\s\"']{6,})",
    re.IGNORECASE,
)
_SECRET_TOKEN = re.compile(
    r"\b(sk-[A-Za-z0-9_-]{12,}|AQ\.[A-Za-z0-9_-]{20,}|AIza[A-Za-z0-9_-]{20,}|"
    r"apikey_[A-Za-z0-9_]{20,}|gh[pousr]_[A-Za-z0-9]{20,})"
)


def _fallback_path_is_sensitive(path: Any) -> bool:
    try:
        resolved = Path(path).expanduser().resolve()
    except (OSError, RuntimeError):
        return True
    lowered = str(resolved).lower()
    if resolved.name.lower() in _FALLBACK_SENSITIVE_NAMES:
        return True
    if resolved.name.startswith(".") and resolved.parent == Path.home():
        return True
    if any(part in lowered for part in _FALLBACK_SENSITIVE_DIRS):
        return True
    return resolved.suffix.lower() in _FALLBACK_SENSITIVE_SUFFIXES


def _fallback_redact_secrets(text: str) -> str:
    if not text:
        return text
    text = _SECRET_ASSIGN.sub(lambda m: m.group(1) + "***REDACTED***", text)
    return _SECRET_TOKEN.sub("***REDACTED***", text)


path_is_sensitive: Callable[[Any], bool]
redact_secrets: Callable[[str], str]
try:  # nexus.safety boshqa agent tomonidan yozilmoqda — bo'lsa undan foydalanamiz
    from nexus.safety import path_is_sensitive, redact_secrets  # type: ignore
except ImportError:  # pragma: no cover - fallback
    path_is_sensitive = _fallback_path_is_sensitive
    redact_secrets = _fallback_redact_secrets

# ---------------------------------------------------------------------------
# Yo'llar
# ---------------------------------------------------------------------------
HOME = Path.home()
DEFAULT_DIR = Path(os.getenv("NEXUS_DEFAULT_DIR", "~/Desktop")).expanduser()
NOTES_FILE = HOME / "Documents" / "nexus_notes.txt"
SPREADSHEET_FILE = HOME / "Documents" / "nexus.xlsx"
IS_WINDOWS = sys.platform == "win32"
if IS_WINDOWS:
    ALLOWED_ROOTS: tuple[Path, ...] = (HOME, Path(os.environ.get("TEMP") or HOME))
else:
    ALLOWED_ROOTS = (HOME, Path("/tmp"), Path("/private/tmp"), Path("/Volumes"))

FOLDER_ALIASES: dict[str, Path] = {
    "desktop": HOME / "Desktop",
    "ish stoli": HOME / "Desktop",
    "ishstoli": HOME / "Desktop",
    "рабочий стол": HOME / "Desktop",
    "documents": HOME / "Documents",
    "hujjatlar": HOME / "Documents",
    "документы": HOME / "Documents",
    "downloads": HOME / "Downloads",
    "yuklamalar": HOME / "Downloads",
    "yuklanmalar": HOME / "Downloads",
    "загрузки": HOME / "Downloads",
    "home": HOME,
    "uy": HOME,
}

MAX_LIST = 100
MAX_READ_CHARS = 4000
_CELL_RE = re.compile(r"^[A-Za-z]{1,3}[1-9][0-9]{0,6}$")
_NUM_RE = re.compile(r"^-?(0|[1-9]\d*)(\.\d+)?$")


def resolve_path(raw: str | None, default: Path | None = None) -> Path:
    """Model aytgan yo'lni absolyut `Path`ga aylantiradi.

    Bo'sh bo'lsa `default` (yoki DEFAULT_DIR). Alias ("ish stoli/hisobot.txt") va `~`,
    `$VAR` qo'llanadi. Nisbiy yo'l DEFAULT_DIR ga nisbatan olinadi.
    """
    if raw is None or not str(raw).strip():
        return default if default is not None else DEFAULT_DIR
    text = str(raw).strip().strip('"').strip("'")
    lowered = text.lower().replace("\\", "/")
    for alias, folder in FOLDER_ALIASES.items():
        if lowered == alias or lowered == alias + "/":
            return folder
        prefix = alias + "/"
        if lowered.startswith(prefix):
            return folder / text[len(prefix) :]
        prefix = "~/" + alias + "/"
        if lowered.startswith(prefix):
            return folder / text[len(prefix) :]
    path = Path(os.path.expandvars(os.path.expanduser(text)))
    if not path.is_absolute():
        path = DEFAULT_DIR / path
    return path


def path_allowed(path: Path) -> tuple[bool, str]:
    """Yo'l ruxsat etilgan ildizlar ostida va sezgir emasligini tekshiradi."""
    try:
        resolved = path.expanduser().resolve()
    except (OSError, RuntimeError) as e:
        return False, f"Yo'lni aniqlab bo'lmadi: {e}"
    roots = [r.resolve() if r.exists() else r for r in ALLOWED_ROOTS]
    if not any(resolved == r or r in resolved.parents for r in roots):
        return False, f"Ruxsat etilmagan yo'l: {resolved} (faqat uy papkasi, /tmp va /Volumes)"
    if path_is_sensitive(resolved):
        return False, f"Sezgir yo'l rad etildi: {resolved}"
    return True, ""


def _result(ok: bool, output: Any, **extra: Any) -> dict[str, Any]:
    text = output if isinstance(output, str) else json.dumps(output, ensure_ascii=False)
    d: dict[str, Any] = {"ok": ok, "output": text}
    d.update(extra)
    return d


def _is_confirmed(args: dict[str, Any]) -> bool:
    v = args.get("confirmed")
    if isinstance(v, str):
        return v.strip().lower() in {"1", "true", "yes", "ha"}
    return bool(v)


def _coerce_cell_value(value: Any) -> Any:
    """"250000" → 250000, "3.5" → 3.5; qolgani satr. Telefon (+998...) va 007 satr qoladi."""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = "" if value is None else str(value)
    stripped = text.strip()
    if _NUM_RE.match(stripped):
        try:
            return int(stripped) if "." not in stripped else float(stripped)
        except ValueError:
            return text
    return text


def _to_list(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, str):
        # "Ali, Toshkent, 250" ko'rinishidagi satr ham qabul qilinadi
        return [p.strip() for p in value.split(",")] if "," in value else [value]
    return ["" if v is None else str(v) for v in value]


# ---------------------------------------------------------------------------
# Oddiy fayllar
# ---------------------------------------------------------------------------
def _read_file_sync(path: Path, max_chars: int) -> dict[str, Any]:
    ok, why = path_allowed(path)
    if not ok:
        return _result(False, why)
    if not path.is_file():
        return _result(False, f"Fayl topilmadi: {path}")
    try:
        with path.open("r", encoding="utf-8", errors="replace") as fh:
            content = fh.read(max_chars + 1)
    except OSError as e:
        return _result(False, f"O'qib bo'lmadi: {e}")
    truncated = len(content) > max_chars
    content = redact_secrets(content[:max_chars])
    out = {"path": str(path), "truncated": truncated, "content": content}
    note = f"Fayl {max_chars} belgidan uzun — qisqartirildi." if truncated else None
    return _result(True, out, **({"note": note} if note else {}))


async def read_file(args: dict[str, Any]) -> dict[str, Any]:
    path = resolve_path(args.get("path"))
    max_chars = int(args.get("max_chars") or MAX_READ_CHARS)
    max_chars = max(100, min(max_chars, 20000))
    return await asyncio.to_thread(_read_file_sync, path, max_chars)


def _write_file_sync(path: Path, content: str, mode: str, confirmed: bool) -> dict[str, Any]:
    ok, why = path_allowed(path)
    if not ok:
        return _result(False, why)
    if path.is_dir():
        return _result(False, f"Bu papka, fayl emas: {path}")
    exists = path.is_file()
    if mode == "overwrite" and exists and not confirmed:
        summary = f"Mavjud faylni qayta yozish: {path} ({path.stat().st_size} bayt)"
        return _result(False, summary, needs_confirmation=True, summary=summary)
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        text = content if content.endswith("\n") else content + "\n"
        with path.open("a" if mode == "append" else "w", encoding="utf-8") as fh:
            fh.write(text)
    except OSError as e:
        return _result(False, f"Yozib bo'lmadi: {e}")
    verb = "qo'shildi" if mode == "append" else ("qayta yozildi" if exists else "yaratildi")
    return _result(True, f"{path.name} {verb} ({len(content)} belgi): {path}", path=str(path))


async def write_file(args: dict[str, Any]) -> dict[str, Any]:
    path = resolve_path(args.get("path"))
    content = "" if args.get("content") is None else str(args["content"])
    mode = str(args.get("mode") or "overwrite").lower()
    if mode not in ("overwrite", "append"):
        return _result(False, "mode 'overwrite' yoki 'append' bo'lishi kerak")
    return await asyncio.to_thread(_write_file_sync, path, content, mode, _is_confirmed(args))


async def delete_file(args: dict[str, Any]) -> dict[str, Any]:
    path = resolve_path(args.get("path"))
    ok, why = path_allowed(path)
    if not ok:
        return _result(False, why)
    if not path.exists():
        return _result(False, f"Fayl topilmadi: {path}")
    if path.is_dir():
        return _result(False, "Bu tool faqat bitta faylni o'chiradi; papkani Finder'da o'chiring")
    if not _is_confirmed(args):
        summary = f"Faylni Savatga ko'chirish: {path}"
        return _result(False, summary, needs_confirmation=True, summary=summary)
    if IS_WINDOWS:
        from nexus.windows_actions import _ps_quote, run_powershell

        ok, out = await run_powershell(
            "Add-Type -AssemblyName Microsoft.VisualBasic;"
            "[Microsoft.VisualBasic.FileIO.FileSystem]::DeleteFile("
            f"{_ps_quote(str(path.resolve()))},'OnlyErrorDialogs','SendToRecycleBin')",
            timeout=15,
        )
        if ok:
            return _result(True, f"{path.name} Savatga ko'chirildi: {path}")
        return _result(False, f"Savatga ko'chirib bo'lmadi: {out}")
    script = f"tell application \"Finder\" to delete POSIX file {_as_str(str(path.resolve()))}"
    ok, out = await run_applescript(script, timeout=15)
    if ok:
        return _result(True, f"{path.name} Savatga ko'chirildi: {path}")
    hint = ""
    if "-1743" in out or "not allowed" in out.lower():
        hint = " Automation ruxsati kerak: Privacy & Security → Automation → Finder."
    return _result(False, f"Savatga ko'chirib bo'lmadi: {out}.{hint}")


def _list_directory_sync(folder: Path) -> dict[str, Any]:
    ok, why = path_allowed(folder)
    if not ok:
        return _result(False, why)
    if not folder.is_dir():
        return _result(False, f"Papka topilmadi: {folder}")
    try:
        entries = sorted(folder.iterdir(), key=lambda p: (p.is_file(), p.name.lower()))
    except OSError as e:
        return _result(False, f"O'qib bo'lmadi: {e}")
    entries = [e for e in entries if not e.name.startswith(".")]
    total = len(entries)
    names = [f"{e.name}/" if e.is_dir() else e.name for e in entries[:MAX_LIST]]
    out = {"path": str(folder), "count": total, "entries": names}
    extra = {"note": f"{total} ta element, faqat {MAX_LIST} tasi ko'rsatildi."} if total > MAX_LIST else {}
    return _result(True, out, **extra)


async def list_directory(args: dict[str, Any]) -> dict[str, Any]:
    folder = resolve_path(args.get("path"), DEFAULT_DIR)
    return await asyncio.to_thread(_list_directory_sync, folder)


def _write_note_sync(text: str, path: Path) -> dict[str, Any]:
    ok, why = path_allowed(path)
    if not ok:
        return _result(False, why)
    if path.suffix == "":
        path = path.with_suffix(".txt")
    stamp = datetime.now().astimezone().strftime("%Y-%m-%d %H:%M")
    line = f"[{stamp}] {' '.join(text.split())}\n"
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("a", encoding="utf-8") as fh:
            fh.write(line)
    except OSError as e:
        return _result(False, f"Eslatma yozilmadi: {e}")
    return _result(True, f"Eslatma yozildi: {path.name}", path=str(path), line=line.rstrip("\n"))


async def write_note(args: dict[str, Any]) -> dict[str, Any]:
    text = str(args.get("text") or "").strip()
    if not text:
        return _result(False, "Eslatma matni bo'sh")
    path = resolve_path(args.get("file"), NOTES_FILE)
    return await asyncio.to_thread(_write_note_sync, text, path)


# ---------------------------------------------------------------------------
# Jadval (xlsx / csv)
# ---------------------------------------------------------------------------
def _spreadsheet_path(raw: Any, allow_csv: bool = True) -> Path:
    path = resolve_path(raw, SPREADSHEET_FILE)
    suffix = path.suffix.lower()
    if suffix not in (".xlsx", ".csv") or (suffix == ".csv" and not allow_csv):
        path = path.with_suffix(".xlsx")
    return path


def _append_csv(path: Path, values: list[str], headers: list[str]) -> dict[str, Any]:
    new_file = not path.exists()
    with path.open("a", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        if new_file and headers:
            writer.writerow(headers)
        writer.writerow(values)
    with path.open("r", encoding="utf-8-sig") as fh:
        rows = sum(1 for _ in fh)
    return _result(True, f"{path.name} fayliga {rows}-qator yozildi: {' | '.join(values)}", path=str(path), row=rows)


def _open_book(path: Path, sheet: str | None, headers: list[str]) -> tuple[Any, Any]:
    from openpyxl import Workbook, load_workbook

    if path.exists():
        book = load_workbook(path)
        if sheet and sheet in book.sheetnames:
            page = book[sheet]
        elif sheet:
            page = book.create_sheet(sheet)
            if headers:
                page.append(headers)
        else:
            page = book.active
    else:
        book = Workbook()
        page = book.active
        if sheet:
            page.title = sheet
        if headers:
            page.append(headers)
    return book, page


def _append_xlsx(path: Path, values: list[str], sheet: str | None, headers: list[str]) -> dict[str, Any]:
    book, page = _open_book(path, sheet, headers)
    page.append([_coerce_cell_value(v) for v in values])
    book.save(path)
    return _result(
        True,
        f"{path.name} ({page.title}) {page.max_row}-qator: {' | '.join(values)}",
        path=str(path),
        sheet=page.title,
        row=page.max_row,
    )


def _append_row_sync(path: Path, values: list[str], sheet: str | None, headers: list[str]) -> dict[str, Any]:
    ok, why = path_allowed(path)
    if not ok:
        return _result(False, why)
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        if path.suffix.lower() == ".csv":
            return _append_csv(path, values, headers)
        return _append_xlsx(path, values, sheet, headers)
    except ImportError:
        return _result(False, "openpyxl o'rnatilmagan: uv pip install openpyxl")
    except (OSError, ValueError) as e:
        return _result(False, f"Jadvalga yozib bo'lmadi: {e}")


async def append_spreadsheet_row(args: dict[str, Any]) -> dict[str, Any]:
    values = _to_list(args.get("values"))
    if not values:
        return _result(False, "values bo'sh")
    path = _spreadsheet_path(args.get("file"))
    sheet = str(args.get("sheet")).strip() if args.get("sheet") else None
    headers = _to_list(args.get("headers"))
    return await asyncio.to_thread(_append_row_sync, path, values, sheet, headers)


def _set_cell_sync(path: Path, cell: str, value: Any, sheet: str | None) -> dict[str, Any]:
    ok, why = path_allowed(path)
    if not ok:
        return _result(False, why)
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        book, page = _open_book(path, sheet, [])
        page[cell] = _coerce_cell_value(value)
        book.save(path)
    except ImportError:
        return _result(False, "openpyxl o'rnatilmagan: uv pip install openpyxl")
    except (OSError, ValueError, KeyError) as e:
        return _result(False, f"Katakka yozib bo'lmadi: {e}")
    return _result(True, f"{path.name} ({page.title}) {cell} = {value}", path=str(path), sheet=page.title, cell=cell)


async def set_spreadsheet_cell(args: dict[str, Any]) -> dict[str, Any]:
    cell = str(args.get("cell") or "").strip().upper().replace(" ", "")
    if not _CELL_RE.match(cell):
        return _result(False, f"Katak manzili noto'g'ri: '{cell}' (masalan B7)")
    value = args.get("value")
    path = _spreadsheet_path(args.get("file"), allow_csv=False)
    sheet = str(args.get("sheet")).strip() if args.get("sheet") else None
    return await asyncio.to_thread(_set_cell_sync, path, cell, value, sheet)


def _read_spreadsheet_sync(path: Path, sheet: str | None, max_rows: int) -> dict[str, Any]:
    ok, why = path_allowed(path)
    if not ok:
        return _result(False, why)
    if not path.is_file():
        return _result(False, f"Jadval topilmadi: {path}")
    rows: list[list[str]] = []
    title = None
    sheets: list[str] = []
    try:
        if path.suffix.lower() == ".csv":
            with path.open("r", encoding="utf-8-sig", newline="") as fh:
                for i, row in enumerate(csv.reader(fh)):
                    if i >= max_rows:
                        break
                    rows.append([str(c) for c in row])
        else:
            from openpyxl import load_workbook

            book = load_workbook(path, read_only=True, data_only=True)
            sheets = list(book.sheetnames)
            page = book[sheet] if sheet and sheet in book.sheetnames else book.active
            title = page.title
            for i, row in enumerate(page.iter_rows(values_only=True)):
                if i >= max_rows:
                    break
                rows.append(["" if c is None else str(c) for c in row])
            book.close()
    except ImportError:
        return _result(False, "openpyxl o'rnatilmagan: uv pip install openpyxl")
    except (OSError, ValueError) as e:
        return _result(False, f"Jadvalni o'qib bo'lmadi: {e}")
    lines = [f"{i + 1}: {' | '.join(r)}" for i, r in enumerate(rows)]
    out: dict[str, Any] = {"path": str(path), "rows": len(rows), "text": "\n".join(lines) or "(bo'sh)"}
    if title:
        out["sheet"] = title
    if sheets:
        out["sheets"] = sheets
    return _result(True, out)


async def read_spreadsheet(args: dict[str, Any]) -> dict[str, Any]:
    path = _spreadsheet_path(args.get("file"))
    sheet = str(args.get("sheet")).strip() if args.get("sheet") else None
    max_rows = max(1, min(int(args.get("max_rows") or 25), 500))
    return await asyncio.to_thread(_read_spreadsheet_sync, path, sheet, max_rows)


# ---------------------------------------------------------------------------
# Qidiruv va ochish
# ---------------------------------------------------------------------------
async def find_files(args: dict[str, Any]) -> dict[str, Any]:
    query = str(args.get("query") or "").strip()
    if not query:
        return _result(False, "query bo'sh")
    folder = resolve_path(args.get("folder"), HOME)
    ok, why = path_allowed(folder)
    if not ok:
        return _result(False, why)
    if not folder.is_dir():
        return _result(False, f"Papka topilmadi: {folder}")
    limit = max(1, min(int(args.get("max_results") or 20), 100))

    if IS_WINDOWS:
        lines = await asyncio.to_thread(_walk_find, folder, query, limit * 5)
    else:
        ok, out = await run_shell(["mdfind", "-onlyin", str(folder), "-name", query], timeout=15)
        lines = [ln for ln in out.splitlines() if ln.strip()] if ok else []
    if not lines and not IS_WINDOWS:
        ok2, out2 = await run_shell(["mdfind", "-onlyin", str(folder), query], timeout=15)
        lines = [ln for ln in out2.splitlines() if ln.strip()] if ok2 else []
    results = []
    for ln in lines:
        p = Path(ln)
        allowed, _ = path_allowed(p)
        if allowed:
            results.append(str(p))
        if len(results) >= limit:
            break
    if not results:
        return _result(True, {"query": query, "folder": str(folder), "results": []}, note="Hech narsa topilmadi.")
    return _result(True, {"query": query, "folder": str(folder), "count": len(results), "results": results})


def _walk_find(folder: Path, query: str, limit: int, max_scanned: int = 200_000) -> list[str]:
    """Spotlight (mdfind) yo'q joyda: nomida `query` bor fayllar (yashirin papkalarsiz, cheklangan)."""
    needle = query.lower()
    found: list[str] = []
    scanned = 0
    for root, dirs, files in os.walk(folder):
        dirs[:] = [d for d in dirs if not d.startswith(".") and d.lower() not in ("appdata", "node_modules")]
        for name in dirs + files:
            scanned += 1
            if needle in name.lower():
                found.append(os.path.join(root, name))
                if len(found) >= limit:
                    return found
        if scanned >= max_scanned:
            break
    return found


async def open_with_app(args: dict[str, Any]) -> dict[str, Any]:
    path = resolve_path(args.get("path"))
    app = str(args.get("app") or "").strip()
    if not app:
        return _result(False, "app nomi kerak")
    ok, why = path_allowed(path)
    if not ok:
        return _result(False, why)
    if not path.exists():
        return _result(False, f"Topilmadi: {path}")
    if IS_WINDOWS:
        ok, out = await run_shell(["cmd", "/c", "start", "", app, str(path)], timeout=15)
    else:
        ok, out = await run_shell(["open", "-a", app, str(path)], timeout=15)
    if ok:
        return _result(True, f"{path.name} {app} bilan ochildi")
    return _result(False, "Ochib bo'lmadi: " + (out or "noma'lum xato"))


# ---------------------------------------------------------------------------
# Deklaratsiyalar
# ---------------------------------------------------------------------------
_LANG = " The user may speak Uzbek, Russian or English."
# `confirmed` ATAYIN sxemada e'lon qilinmagan: registry validatsiyasi noma'lum argumentlarni
# tashlab yuboradi, shuning uchun model uni o'zi qo'ya olmaydi — faqat gate tasdiqdan keyin qo'shadi.


def _decl(name: str, description: str, props: dict | None = None, required: list[str] | None = None) -> dict:
    d: dict[str, Any] = {"name": name, "description": description + _LANG}
    if props:
        d["parameters"] = {"type": "OBJECT", "properties": props, "required": required or []}
    return d


def _s(desc: str, enum: list[str] | None = None) -> dict:
    p: dict[str, Any] = {"type": "STRING", "description": desc}
    if enum:
        p["enum"] = enum
    return p


def _i(desc: str) -> dict:
    return {"type": "INTEGER", "description": desc}


def _arr(desc: str) -> dict:
    return {"type": "ARRAY", "items": {"type": "STRING"}, "description": desc}


_PATH_DESC = (
    "File path. Accepts '~', absolute paths, or folder aliases: 'desktop'/'ish stoli', "
    "'downloads'/'yuklamalar', 'documents'/'hujjatlar' (e.g. 'desktop/hisobot.txt'). "
    "A bare file name goes to the Desktop."
)
_SHEET_FILE_DESC = "Spreadsheet path (.xlsx or .csv). Default: ~/Documents/nexus.xlsx."

TOOL_DECLARATIONS: list[dict] = [
    _decl(
        "read_file",
        "Read a text file's content (secrets are redacted). Only files under the home folder are allowed.",
        {"path": _s(_PATH_DESC), "max_chars": _i("Maximum characters to return (default 4000).")},
        ["path"],
    ),
    _decl(
        "write_file",
        "Create or write a text file. Overwriting an existing file requires the user's confirmation "
        "(the system asks automatically). Use mode 'append' to add to the end.",
        {
            "path": _s(_PATH_DESC),
            "content": _s("Text to write."),
            "mode": _s("'overwrite' (default) or 'append'.", ["overwrite", "append"]),
        },
        ["path", "content"],
    ),
    _decl(
        "delete_file",
        "Move a single file to the Trash (recoverable). Always requires the user's confirmation.",
        {"path": _s(_PATH_DESC)},
        ["path"],
    ),
    _decl(
        "list_directory",
        "List the files and folders in a directory (default: Desktop). Folders end with '/'.",
        {"path": _s("Folder path or alias ('ish stoli', 'downloads', 'documents'). Default Desktop.")},
    ),
    _decl(
        "write_note",
        "Append a timestamped note line to the notes file (~/Documents/nexus_notes.txt by default). "
        "Use for 'shuni yozib qo'y', 'eslatma qil', 'запиши'.",
        {"text": _s("Note text."), "file": _s("Optional notes file path.")},
        ["text"],
    ),
    _decl(
        "append_spreadsheet_row",
        "Append one row of values to an Excel (.xlsx) or CSV file, creating it if needed. "
        "Use for 'Excelga yoz: Ali, Toshkent, 250 ming'.",
        {
            "values": _arr("Cell values in column order, e.g. ['Ali', 'Toshkent', '250000']."),
            "file": _s(_SHEET_FILE_DESC),
            "sheet": _s("Sheet name (xlsx only); default active sheet."),
            "headers": _arr("Header row to write when the file/sheet is new."),
        },
        ["values"],
    ),
    _decl(
        "set_spreadsheet_cell",
        "Write a value into a specific cell of an Excel file, e.g. cell 'B7'.",
        {
            "cell": _s("Cell address like 'B7'."),
            "value": _s("Value to write (numbers are stored as numbers)."),
            "file": _s("Excel path (.xlsx). Default: ~/Documents/nexus.xlsx."),
            "sheet": _s("Sheet name; default active sheet."),
        },
        ["cell", "value"],
    ),
    _decl(
        "read_spreadsheet",
        "Read the first rows of an Excel or CSV file to answer 'jadvalda nima bor?'.",
        {
            "file": _s(_SHEET_FILE_DESC),
            "sheet": _s("Sheet name (xlsx only)."),
            "max_rows": _i("Rows to read (default 25)."),
        },
    ),
    _decl(
        "find_files",
        "Search files by name (Spotlight) under a folder, e.g. 'hisobot' in Documents.",
        {
            "query": _s("File name or keyword."),
            "folder": _s("Folder to search in (alias allowed). Default: home folder."),
            "max_results": _i("Maximum results (default 20)."),
        },
        ["query"],
    ),
    _decl(
        "open_with_app",
        "Open a file with a specific application, e.g. open 'hisobot.xlsx' with 'Numbers' or 'Microsoft Excel'.",
        {"path": _s(_PATH_DESC), "app": _s("Application name, e.g. 'TextEdit', 'Preview', 'Numbers'.")},
        ["path", "app"],
    ),
]

HANDLERS: dict[str, Any] = {
    "read_file": read_file,
    "write_file": write_file,
    "delete_file": delete_file,
    "list_directory": list_directory,
    "write_note": write_note,
    "append_spreadsheet_row": append_spreadsheet_row,
    "set_spreadsheet_cell": set_spreadsheet_cell,
    "read_spreadsheet": read_spreadsheet,
    "find_files": find_files,
    "open_with_app": open_with_app,
}

__all__ = [
    "ALLOWED_ROOTS",
    "DEFAULT_DIR",
    "FOLDER_ALIASES",
    "HANDLERS",
    "NOTES_FILE",
    "SPREADSHEET_FILE",
    "TOOL_DECLARATIONS",
    "path_allowed",
    "path_is_sensitive",
    "redact_secrets",
    "resolve_path",
]
