"""Kengaytma modullari (ax_actions, file_actions) testlari — real AX/Finder/Spotlight'ga tegmaydi."""

from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest

from nexus import ax_actions as ax
from nexus import file_actions as fa

VALID_TYPES = {"OBJECT", "STRING", "INTEGER", "NUMBER", "BOOLEAN", "ARRAY"}


# ---------------------------------------------------------------------------
# Kontrakt: TOOL_DECLARATIONS <-> HANDLERS
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("mod", [ax, fa], ids=["ax_actions", "file_actions"])
def test_declarations_match_handlers(mod):
    names = [d["name"] for d in mod.TOOL_DECLARATIONS]
    assert len(names) == len(set(names)), "takroriy tool nomi"
    assert set(names) == set(mod.HANDLERS), "deklaratsiya va handler ro'yxati mos emas"
    for d in mod.TOOL_DECLARATIONS:
        assert d["description"]
        params = d.get("parameters")
        if params is not None:
            assert params["type"] == "OBJECT"
            assert params["properties"], "bo'sh OBJECT bo'lmasin"
            for key, spec in params["properties"].items():
                assert spec["type"] in VALID_TYPES, f"{d['name']}.{key}: {spec['type']}"
                if spec["type"] == "ARRAY":
                    assert spec["items"]["type"] in VALID_TYPES
            for req in params.get("required", []):
                assert req in params["properties"]
        json.dumps(d)  # seriyalanadi


def test_declaration_names_unique_across_modules():
    from nexus.tools.schemas import ALL_TOOL_DECLARATIONS

    core = {d["name"] for d in ALL_TOOL_DECLARATIONS}
    ext = [d["name"] for d in ax.TOOL_DECLARATIONS + fa.TOOL_DECLARATIONS]
    assert len(ext) == len(set(ext))
    assert not core & set(ext), "asosiy sxemalar bilan nom to'qnashuvi"


def test_confirmed_not_declared():
    """`confirmed` sxemada bo'lmasligi shart — aks holda model gate'ni chetlab o'tadi."""
    for d in fa.TOOL_DECLARATIONS:
        props = (d.get("parameters") or {}).get("properties") or {}
        assert "confirmed" not in props, d["name"]


@pytest.mark.parametrize("mod", [ax, fa], ids=["ax_actions", "file_actions"])
async def test_handlers_are_async_and_return_contract(mod):
    import inspect

    for name, h in mod.HANDLERS.items():
        assert inspect.iscoroutinefunction(h), name


# ---------------------------------------------------------------------------
# file_actions: yo'llar
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    "raw,expected",
    [
        ("desktop", Path.home() / "Desktop"),
        ("Ish stoli", Path.home() / "Desktop"),
        ("рабочий стол", Path.home() / "Desktop"),
        ("downloads/a.pdf", Path.home() / "Downloads" / "a.pdf"),
        ("yuklamalar", Path.home() / "Downloads"),
        ("загрузки", Path.home() / "Downloads"),
        ("hujjatlar/hisobot.xlsx", Path.home() / "Documents" / "hisobot.xlsx"),
        ("documents", Path.home() / "Documents"),
        ("home", Path.home()),
        ("~/Music", Path.home() / "Music"),
        ('"~/Music"', Path.home() / "Music"),
        ("/tmp/x.txt", Path("/tmp/x.txt")),
    ],
)
def test_resolve_path_aliases(raw, expected):
    assert fa.resolve_path(raw) == expected


def test_resolve_path_relative_and_default(monkeypatch, tmp_path):
    monkeypatch.setattr(fa, "DEFAULT_DIR", tmp_path)
    assert fa.resolve_path("hisobot.txt") == tmp_path / "hisobot.txt"
    assert fa.resolve_path(None) == tmp_path
    assert fa.resolve_path("", default=Path("/tmp/n.txt")) == Path("/tmp/n.txt")


@pytest.mark.parametrize(
    "path",
    [
        "~/.zshrc",
        "~/.ssh/id_rsa",
        "~/.ssh/config",
        "~/Library/LaunchAgents/x.plist",
        "~/.env",
        "~/Desktop/run.sh",
    ],
)
def test_sensitive_paths_rejected(path):
    ok, why = fa.path_allowed(Path(path).expanduser())
    assert not ok
    assert "Sezgir" in why


@pytest.mark.parametrize("path", ["/etc/hosts", "/usr/bin/python3", "/System/Library/x.txt", "/var/db/x"])
def test_outside_home_rejected(path):
    ok, _why = fa.path_allowed(Path(path))
    assert not ok


def test_allowed_paths():
    assert fa.path_allowed(Path.home() / "Desktop" / "hisobot.txt")[0]
    assert fa.path_allowed(Path("/tmp/nexus_test.txt"))[0]


@pytest.fixture
def sandbox(monkeypatch, tmp_path):
    """tmp_path ni ruxsat etilgan ildiz va default papka qiladi."""
    monkeypatch.setattr(fa, "ALLOWED_ROOTS", (tmp_path,))
    monkeypatch.setattr(fa, "DEFAULT_DIR", tmp_path)
    monkeypatch.setattr(fa, "NOTES_FILE", tmp_path / "notes" / "nexus_notes.txt")
    monkeypatch.setattr(fa, "SPREADSHEET_FILE", tmp_path / "nexus.xlsx")
    return tmp_path


# ---------------------------------------------------------------------------
# write_file / read_file / delete_file
# ---------------------------------------------------------------------------
async def test_write_file_new_then_needs_confirmation_then_confirmed(sandbox):
    target = sandbox / "hisobot.txt"
    res = await fa.HANDLERS["write_file"]({"path": str(target), "content": "salom"})
    assert res["ok"] is True
    assert target.read_text(encoding="utf-8") == "salom\n"

    res = await fa.HANDLERS["write_file"]({"path": str(target), "content": "yangi"})
    assert res["ok"] is False
    assert res["needs_confirmation"] is True
    assert "hisobot.txt" in res["summary"]
    assert target.read_text(encoding="utf-8") == "salom\n", "tasdiqsiz yozilmasligi kerak"

    res = await fa.HANDLERS["write_file"]({"path": str(target), "content": "yangi", "confirmed": True})
    assert res["ok"] is True
    assert target.read_text(encoding="utf-8") == "yangi\n"


async def test_write_file_append_needs_no_confirmation(sandbox):
    target = sandbox / "a.txt"
    await fa.HANDLERS["write_file"]({"path": str(target), "content": "1"})
    res = await fa.HANDLERS["write_file"]({"path": str(target), "content": "2", "mode": "append"})
    assert res["ok"] is True
    assert target.read_text() == "1\n2\n"


async def test_write_file_bad_mode_and_sensitive(sandbox, monkeypatch):
    res = await fa.HANDLERS["write_file"]({"path": "x.txt", "mode": "zzz", "content": "a"})
    assert res["ok"] is False
    monkeypatch.setattr(fa, "ALLOWED_ROOTS", (Path.home(),))
    res = await fa.HANDLERS["write_file"]({"path": "~/.zshrc", "content": "evil", "mode": "append"})
    assert res["ok"] is False and "Sezgir" in res["output"]


async def test_read_file_truncation_and_redaction(sandbox):
    target = sandbox / "big.txt"
    target.write_text("API_KEY=abcdef123456\n" + "x" * 500, encoding="utf-8")
    res = await fa.HANDLERS["read_file"]({"path": "big.txt", "max_chars": 100})
    assert res["ok"] is True
    data = json.loads(res["output"])
    assert data["truncated"] is True
    assert data["content"].count("x") < 100  # 100 belgi + redaksiya farqi
    assert "abcdef123456" not in data["content"]
    assert "REDACTED" in data["content"]
    assert "qisqartirildi" in res["note"]

    res = await fa.HANDLERS["read_file"]({"path": "yoq.txt"})
    assert res["ok"] is False and "topilmadi" in res["output"]


async def test_delete_file_moves_to_trash_via_finder(sandbox, monkeypatch):
    target = sandbox / "del.txt"
    target.write_text("x")
    res = await fa.HANDLERS["delete_file"]({"path": str(target)})
    assert res["ok"] is False and res["needs_confirmation"] is True
    assert target.exists()

    scripts: list[str] = []

    async def fake_applescript(script: str, timeout: float = 15) -> tuple[bool, str]:
        scripts.append(script)
        return True, ""

    monkeypatch.setattr(fa, "run_applescript", fake_applescript)
    res = await fa.HANDLERS["delete_file"]({"path": str(target), "confirmed": True})
    assert res["ok"] is True
    assert scripts and 'tell application "Finder" to delete POSIX file' in scripts[0]
    assert str(target.resolve()) in scripts[0]
    assert target.exists(), "python o'zi o'chirmasligi kerak — faqat Finder Savatga ko'chiradi"

    res = await fa.HANDLERS["delete_file"]({"path": str(sandbox), "confirmed": True})
    assert res["ok"] is False and "papka" in res["output"].lower()


# ---------------------------------------------------------------------------
# list_directory / write_note
# ---------------------------------------------------------------------------
async def test_list_directory(sandbox):
    (sandbox / "b.txt").write_text("b")
    (sandbox / "sub").mkdir()
    (sandbox / ".hidden").write_text("h")
    res = await fa.HANDLERS["list_directory"]({})
    data = json.loads(res["output"])
    assert res["ok"] and data["entries"] == ["sub/", "b.txt"]
    assert data["count"] == 2

    for i in range(fa.MAX_LIST + 5):
        (sandbox / f"f{i:03d}.txt").write_text("")
    res = await fa.HANDLERS["list_directory"]({"path": str(sandbox)})
    data = json.loads(res["output"])
    assert len(data["entries"]) == fa.MAX_LIST
    assert data["count"] == fa.MAX_LIST + 7
    assert "note" in res

    res = await fa.HANDLERS["list_directory"]({"path": str(sandbox / "yoq")})
    assert res["ok"] is False


async def test_write_note_format(sandbox):
    import re

    res = await fa.HANDLERS["write_note"]({"text": "ertaga  soat uchda\nuchrashuv"})
    assert res["ok"] is True
    notes = sandbox / "notes" / "nexus_notes.txt"
    lines = notes.read_text(encoding="utf-8").splitlines()
    assert len(lines) == 1
    assert re.fullmatch(r"\[\d{4}-\d{2}-\d{2} \d{2}:\d{2}\] ertaga soat uchda uchrashuv", lines[0])

    await fa.HANDLERS["write_note"]({"text": "ikkinchi", "file": "custom"})
    assert (sandbox / "custom.txt").read_text().endswith("ikkinchi\n")

    res = await fa.HANDLERS["write_note"]({"text": "   "})
    assert res["ok"] is False


# ---------------------------------------------------------------------------
# Jadval: xlsx va csv
# ---------------------------------------------------------------------------
async def test_append_and_read_xlsx(sandbox):
    from openpyxl import load_workbook

    res = await fa.HANDLERS["append_spreadsheet_row"](
        {"values": ["Ali", "Toshkent", "250000"], "headers": ["Ism", "Shahar", "Summa"]}
    )
    assert res["ok"] is True, res
    assert res["row"] == 2
    res = await fa.HANDLERS["append_spreadsheet_row"]({"values": "Vali, Samarqand, 12.5"})
    assert res["ok"] and res["row"] == 3

    book = load_workbook(sandbox / "nexus.xlsx")
    ws = book.active
    assert [c.value for c in ws[1]] == ["Ism", "Shahar", "Summa"]
    assert ws["C2"].value == 250000  # raqam sifatida
    assert ws["C3"].value == 12.5

    res = await fa.HANDLERS["set_spreadsheet_cell"]({"cell": "b7", "value": "500"})
    assert res["ok"] and res["cell"] == "B7"
    book = load_workbook(sandbox / "nexus.xlsx")
    assert book.active["B7"].value == 500

    res = await fa.HANDLERS["set_spreadsheet_cell"]({"cell": "7B", "value": "1"})
    assert res["ok"] is False

    res = await fa.HANDLERS["read_spreadsheet"]({"max_rows": 2})
    data = json.loads(res["output"])
    assert res["ok"] and data["rows"] == 2
    assert data["text"].splitlines()[0] == "1: Ism | Shahar | Summa"
    assert data["text"].splitlines()[1] == "2: Ali | Toshkent | 250000"


async def test_append_and_read_csv(sandbox):
    csv_path = sandbox / "data.csv"
    res = await fa.HANDLERS["append_spreadsheet_row"](
        {"values": ["a", "b"], "file": str(csv_path), "headers": ["h1", "h2"]}
    )
    assert res["ok"] and res["row"] == 2
    await fa.HANDLERS["append_spreadsheet_row"]({"values": ["c", "d"], "file": "data.csv"})
    raw = csv_path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf"), "utf-8-sig BOM (Excel uchun)"
    with csv_path.open(encoding="utf-8-sig", newline="") as fh:
        assert list(csv.reader(fh)) == [["h1", "h2"], ["a", "b"], ["c", "d"]]

    res = await fa.HANDLERS["read_spreadsheet"]({"file": "data.csv"})
    data = json.loads(res["output"])
    assert data["rows"] == 3 and data["text"].splitlines()[-1] == "3: c | d"

    res = await fa.HANDLERS["read_spreadsheet"]({"file": "yoq.csv"})
    assert res["ok"] is False


def test_coerce_cell_value():
    assert fa._coerce_cell_value("250000") == 250000
    assert fa._coerce_cell_value("3.5") == 3.5
    assert fa._coerce_cell_value("+998901234567") == "+998901234567"
    assert fa._coerce_cell_value("007") == "007"
    assert fa._coerce_cell_value("Ali") == "Ali"
    assert fa._coerce_cell_value(None) == ""


# ---------------------------------------------------------------------------
# ax_actions: toza funksiyalar
# ---------------------------------------------------------------------------
def test_pick_label_priority():
    assert ax.pick_label({"AXTitle": "Save", "AXDescription": "desc"}) == "Save"
    assert ax.pick_label({"AXTitle": "  ", "AXDescription": "Close window"}) == "Close window"
    assert ax.pick_label({"AXTitle": None, "AXValue": 42, "AXHelp": "Help text"}) == "Help text"
    assert ax.pick_label({"AXPlaceholderValue": "Search", "AXRoleDescription": "text field"}) == "Search"
    assert ax.pick_label({"AXRoleDescription": "button"}) == "button"
    assert ax.pick_label({}) == ""
    assert ax.pick_label({"AXTitle": "a   b\n c"}) == "a b c"
    assert len(ax.pick_label({"AXTitle": "x" * 200})) == ax.LABEL_MAX


@pytest.mark.parametrize(
    "wanted,label,expected",
    [
        ("Save", "save", 1.0),
        ("save", "Save As…", 0.85),
        ("new tab", "New Tab", 1.0),
        ("Yangi hujjat", "Hujjat yangi", 0.8),
        ("Reload", "Reloaf", None),  # fuzzy: 0.72..0.99
        ("Sign in", "Sign out", 0.0),  # boshqa amal — mos kelmasligi to'g'ri
        ("Bookmark", "Extensions", 0.0),
        ("", "Save", 0.0),
        ("Save", "", 0.0),
    ],
)
def test_match_label(wanted, label, expected):
    score = ax._match_label(wanted, label)
    if expected is None:
        assert 0.72 <= score < 1.0
    else:
        assert score == expected


def test_rank_matches_prefers_exact_then_bigger():
    els = [
        ax.Element(1, "AXButton", "New Tab", (10, 10), (20, 20)),
        ax.Element(2, "AXButton", "New", (50, 50), (100, 40)),
        ax.Element(3, "AXButton", "New", (90, 90), (20, 20)),
        ax.Element(4, "AXButton", "Extensions", (0, 0), (30, 30)),
    ]
    ranked = ax.rank_matches("new", els)
    assert [e.n for _, e in ranked] == [2, 3, 1]
    assert ranked[0][0] == 1.0
    assert ax.rank_matches("zzz", els) == []


def test_element_as_dict():
    el = ax.Element(3, "AXLink", "Mijozlar", (120, 387), (80, 20))
    assert el.as_dict() == {"n": 3, "role": "Link", "label": "Mijozlar", "at": [120, 387], "size": [80, 20]}


async def test_ax_handlers_without_pyobjc(monkeypatch):
    """pyobjc yo'q holat: handler xato matni qaytaradi, istisno emas."""
    monkeypatch.setattr(ax, "_AS", None)
    monkeypatch.setattr(ax, "_NSWorkspace", None)
    monkeypatch.setattr(ax, "_QZ", None)
    for name, h in ax.HANDLERS.items():
        res = await h({"label": "x", "path": ["File"], "text": "t"})
        assert res["ok"] is False, name
        assert "pyobjc" in res["output"]


async def test_ax_handlers_without_accessibility(monkeypatch):
    monkeypatch.setattr(ax, "is_trusted", lambda: False)
    res = await ax.HANDLERS["list_ui_elements"]({})
    if ax.available():
        assert res["ok"] is False and "Accessibility" in res["output"]
    else:
        assert res["ok"] is False


async def test_click_ui_element_needs_choice_and_index(monkeypatch):
    if not ax.available():
        pytest.skip("pyobjc yo'q")
    monkeypatch.setattr(ax, "is_trusted", lambda: True)
    els = [
        ax.Element(1, "AXButton", "New", (50, 50), (100, 40), ref=object()),
        ax.Element(2, "AXButton", "New", (90, 90), (20, 20), ref=object()),
        ax.Element(3, "AXButton", "Save", (10, 10), (20, 20), ref=object()),
    ]
    monkeypatch.setattr(ax, "collect_elements", lambda limit=120: ("App", els))
    monkeypatch.setattr(ax, "_ocr_read", lambda max_chars=6000: None)  # real ekranga tegmasin
    monkeypatch.setattr(ax, "front_app", lambda: ("App", 1))
    clicked: list[str] = []
    monkeypatch.setattr(
        ax, "activate_element", lambda el: (clicked.append(el.label), (True, f"'{el.label}' bosildi"))[1]
    )
    monkeypatch.setattr(ax, "_geometry", lambda ref: ((0, 0), (1, 1)))

    res = await ax.HANDLERS["click_ui_element"]({"label": "new"})
    assert res["ok"] is False and res["needs_choice"] is True
    assert [m["n"] for m in res["matches"]] == [1, 2]
    assert clicked == []

    res = await ax.HANDLERS["click_ui_element"]({"index": 2})
    assert res["ok"] is True and clicked == ["New"]

    res = await ax.HANDLERS["click_ui_element"]({"label": "save"})
    assert res["ok"] is True and clicked[-1] == "Save"

    res = await ax.HANDLERS["click_ui_element"]({"label": "yo'q tugma"})
    assert res["ok"] is False and "topilmadi" in res["output"]

    res = await ax.HANDLERS["click_ui_element"]({})
    assert res["ok"] is False


async def test_menu_command_path_parsing(monkeypatch):
    if not ax.available():
        pytest.skip("pyobjc yo'q")
    monkeypatch.setattr(ax, "is_trusted", lambda: True)
    seen: list[list[str]] = []
    monkeypatch.setattr(ax, "run_menu", lambda parts: (seen.append(parts), (True, "ok"))[1])
    await ax.HANDLERS["menu_command"]({"path": ["File", "Save"]})
    await ax.HANDLERS["menu_command"]({"path": "File > Save As"})
    assert seen == [["File", "Save"], ["File", "Save As"]]
    res = await ax.HANDLERS["menu_command"]({"path": []})
    assert res["ok"] is False
